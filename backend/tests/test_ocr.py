from io import BytesIO
from pathlib import Path

import pymupdf as fitz
import pytest
from PIL import Image, ImageDraw, ImageFont

from eic.knowledge.documents import extract_text


def _font() -> ImageFont.FreeTypeFont:
    candidates = [
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ]
    font_path = next((path for path in candidates if path.exists()), None)
    if font_path is None:
        pytest.skip("no portable test font is installed")
    return ImageFont.truetype(font_path, 52)


def test_image_ocr_extracts_fixture_text(tmp_path: Path) -> None:
    image_path = tmp_path / "ocr.png"
    image = Image.new("RGB", (1400, 240), "white")
    draw = ImageDraw.Draw(image)
    font = _font()
    draw.text((30, 75), "Enterprise Intelligence Center OCR 123", fill="black", font=font)
    image.save(image_path)

    extracted = " ".join(extract_text(image_path).split())

    assert "Enterprise Intelligence Center" in extracted
    assert "OCR 123" in extracted


def test_scanned_pdf_uses_ocr_fallback(tmp_path: Path) -> None:
    image = Image.new("RGB", (1600, 300), "white")
    draw = ImageDraw.Draw(image)
    font = _font()
    draw.text((40, 100), "Enterprise Intelligence Center SCANNED 456", fill="black", font=font)
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    pdf_path = tmp_path / "scanned.pdf"
    document = fitz.open()
    page = document.new_page(width=800, height=150)
    page.insert_image(page.rect, stream=buffer.getvalue())
    document.save(pdf_path)
    document.close()

    extracted = " ".join(extract_text(pdf_path).split())

    assert "SCANNED 456" in extracted


def test_legacy_doc_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "legacy.doc"
    path.write_bytes(b"not a real word document")
    try:
        extract_text(path)
    except ValueError as exc:
        assert ".doc" in str(exc)
    else:
        raise AssertionError(".doc was accepted")


def test_mixed_pdf_keeps_text_layer_and_ocr_embedded_image(tmp_path: Path) -> None:
    from eic.knowledge.documents import extract_document

    image = Image.new("RGB", (1600, 300), "white")
    draw = ImageDraw.Draw(image)
    font = _font()
    draw.text((40, 100), "EMBEDDED IMAGE TEXT 789", fill="black", font=font)
    buffer = BytesIO()
    image.save(buffer, format="PNG")

    pdf_path = tmp_path / "mixed.pdf"
    document = fitz.open()
    page = document.new_page(width=800, height=420)
    page.insert_text((40, 48), "VISIBLE TEXT LAYER HELLO", fontsize=16)
    page.insert_image(fitz.Rect(40, 90, 760, 280), stream=buffer.getvalue())
    document.save(pdf_path)
    document.close()

    extracted = extract_document(pdf_path)
    blob = " ".join(extracted.text.split())
    assert "VISIBLE TEXT LAYER HELLO" in blob
    assert "EMBEDDED IMAGE TEXT 789" in blob
    assert "text_layer" in extracted.report["methods"]
    assert "embedded_image_ocr" in extracted.report["methods"]
    assert "第1页" in extracted.text


def test_pdf_tables_are_extracted_as_markdown(tmp_path: Path) -> None:
    from eic.knowledge.documents import extract_document

    pdf_path = tmp_path / "table.pdf"
    document = fitz.open()
    page = document.new_page(width=500, height=240)
    page.insert_htmlbox(
        page.rect,
        "<table border='1'><tr><th>HA</th><th>HB</th></tr><tr><td>CELLALPHA</td><td>CELLBETA</td></tr></table>",
    )
    document.save(pdf_path)
    document.close()

    extracted = extract_document(pdf_path)
    blob = extracted.text
    assert "CELLALPHA" in blob
    assert "CELLBETA" in blob
    if "table_markdown" in extracted.report["methods"]:
        assert "|" in blob


def test_docx_tables_and_embedded_images(tmp_path: Path) -> None:
    from docx import Document
    from docx.shared import Inches

    from eic.knowledge.documents import extract_document

    image_path = tmp_path / "word-ocr.png"
    image = Image.new("RGB", (1400, 240), "white")
    draw = ImageDraw.Draw(image)
    font = _font()
    draw.text((30, 75), "WORDOCRTOKEN 321", fill="black", font=font)
    image.save(image_path)

    docx_path = tmp_path / "mixed.docx"
    document = Document()
    document.add_paragraph("WORD PARAGRAPH ALPHA")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "H1"
    table.cell(0, 1).text = "H2"
    table.cell(1, 0).text = "CELLONE"
    table.cell(1, 1).text = "CELLTWO"
    document.add_picture(str(image_path), width=Inches(5))
    document.save(docx_path)

    extracted = extract_document(docx_path)
    blob = " ".join(extracted.text.split())
    assert "WORD PARAGRAPH ALPHA" in blob
    assert "CELLONE" in blob
    assert "CELLTWO" in blob
    assert "WORDOCRTOKEN" in blob
    assert "321" in blob
    assert "table_markdown" in extracted.report["methods"]
    assert "embedded_image_ocr" in extracted.report["methods"]
    assert "[表]" in extracted.text
    assert "[图]" in extracted.text
