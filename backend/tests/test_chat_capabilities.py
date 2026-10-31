"""附件解析、联网来源及向量批次的隔离契约检查。"""
from contextlib import asynccontextmanager
from datetime import datetime, timedelta, UTC
from io import BytesIO
import json
from types import SimpleNamespace
import zipfile

import httpx
import pytest
from pydantic import SecretStr

from eic.ai import attachments, embeddings, search
from eic.ai.budget import ChatConfigurationError
from sqlalchemy.ext.asyncio import AsyncSession, create_async_engine

from eic.models import AIModel, ChatMessage, KnowledgeDocument, Resource


@pytest.fixture
async def attachment_session():
    """仅创建 helper 查询所需的内存表，不加载应用或开发数据库。"""
    engine = create_async_engine('sqlite+aiosqlite:///:memory:')
    try:
        async with engine.begin() as connection:
            for table in (Resource.__table__, ChatMessage.__table__, KnowledgeDocument.__table__):
                await connection.run_sync(table.create)
        async with AsyncSession(engine, expire_on_commit=False) as session:
            yield session
    finally:
        await engine.dispose()


def add_attachment_message(session, identifier, files, *, conversation=10, tenant=1,
                           owner=2, role='user', deleted=False, created_at=None):
    session.add_all([
        Resource(id=identifier, tenant_id=tenant, owner_id=owner,
                 resource_type='chat_messages', data={'fileIds': files}, is_deleted=deleted),
        ChatMessage(id=identifier, tenant_id=tenant, session_id=conversation,
                    role=role, content='fixture', created_at=created_at or datetime(2026, 1, 1, tzinfo=UTC)),
    ])


@pytest.mark.asyncio
async def test_attachment_history_prefers_recent_and_respects_boundary_and_scope(attachment_session):
    session = attachment_session
    principal = SimpleNamespace(tenant_id=1, user_id=2)
    now = datetime(2026, 1, 1, tzinfo=UTC)
    add_attachment_message(session, 100, [11], created_at=now - timedelta(days=1))
    add_attachment_message(session, 90, [12], created_at=now)
    add_attachment_message(session, 91, [13], created_at=now)
    add_attachment_message(session, 92, [14], created_at=now)
    for identifier, options in [(93, {'conversation': 20}), (94, {'tenant': 3}),
                                (95, {'owner': 4}), (96, {'role': 'assistant'}),
                                (97, {'deleted': True})]:
        add_attachment_message(session, identifier, [identifier], created_at=now - timedelta(hours=1), **options)
    await session.flush()
    ids = await attachments.conversation_file_ids([20], principal, session, 10,
                                                 before_message_id=91, max_files=3)
    assert ids == [20, 13, 12]
    assert ids.omitted_history_ids == [11]


@pytest.mark.asyncio
@pytest.mark.parametrize('message_count', [500, 501])
@pytest.mark.parametrize('has_recent_file', [False, True])
async def test_attachment_history_scan_limit_preserves_context(
        attachment_session, monkeypatch, tmp_path, message_count, has_recent_file):
    principal = SimpleNamespace(tenant_id=1, user_id=2)
    for identifier in range(1, message_count + 1):
        files = [9001] if has_recent_file and identifier == message_count - 1 else []
        add_attachment_message(attachment_session, identifier, files)
    await attachment_session.flush()
    # 最新一条消息没有附件；超过扫描窗口不应阻断普通聊天。
    ids = await attachments.conversation_file_ids(
        [], principal, attachment_session, 10,
        before_message_id=message_count, max_files=1)
    assert ids == ([9001] if has_recent_file else [])
    assert ids.history_truncated is (message_count == 501)
    assert ids.omitted_history_ids == []

    source = tmp_path / 'recent.txt'
    source.write_bytes(b'recent attachment fixture')
    reads = []

    async def authorize(session, current_principal, value):
        assert session is attachment_session and current_principal is principal
        reads.append(value)
        return SimpleNamespace(id=value, name='recent.txt')

    monkeypatch.setattr(attachments, 'require_file_access', authorize)
    monkeypatch.setattr(attachments, 'file_source_path', lambda file: source)
    context, snapshots = await attachments.load_context(
        ids, principal, attachment_session,
        limits(chat_attachments_enabled=True, chat_attachment_max_files=1))
    assert reads == ([9001] if has_recent_file else [])
    assert [row['id'] for row in snapshots] == reads
    if has_recent_file:
        assert 'recent attachment fixture' in context
        assert snapshots[0]['size'] == len(source.read_bytes())
    if message_count == 501:
        assert '[附件回读范围说明]' in context
        assert '历史附件仅扫描最近 500 条用户消息' in context
        assert '更早消息未扫描' in context
        assert '不得视为已读取材料' in context
    else:
        assert '[附件回读范围说明]' not in context
        if not has_recent_file:
            assert context == '' and snapshots == []


@pytest.mark.asyncio
async def test_explicit_attachment_limit_rejects_before_history_query():
    class NoQuery:
        async def scalars(self, query):
            pytest.fail('超限的显式附件不应查询历史')
    with pytest.raises(ChatConfigurationError) as error:
        await attachments.conversation_file_ids([1, '2', 3],
            SimpleNamespace(tenant_id=1, user_id=2), NoQuery(), 10, max_files=2)
    assert error.value.reason == 'attachment_budget_exceeded'


@pytest.mark.asyncio
async def test_attachment_ids_deduplicate_int_and_string_before_budget(attachment_session):
    principal = SimpleNamespace(tenant_id=1, user_id=2)
    add_attachment_message(attachment_session, 100, ['1', 2, '2', 3, '3'])
    await attachment_session.flush()
    ids = await attachments.conversation_file_ids([1, '1'], principal, attachment_session,
                                                 10, max_files=2)
    assert ids == [1, 2]
    assert ids.omitted_history_ids == [3]
    explicit = await attachments.conversation_file_ids([1, '1', '2', 2], principal,
                                                      attachment_session, max_files=2)
    assert explicit == [1, 2]


@pytest.mark.asyncio
async def test_load_context_deduplicates_and_reports_omitted_history(monkeypatch, tmp_path):
    source = tmp_path / 'fixture.txt'
    source.write_bytes(b'fixture text')
    reads = []
    async def authorize(session, principal, value):
        reads.append(value)
        return SimpleNamespace(id=value, name='fixture.txt')
    monkeypatch.setattr(attachments, 'require_file_access', authorize)
    monkeypatch.setattr(attachments, 'file_source_path', lambda file: source)
    config = limits(chat_attachments_enabled=True, chat_attachment_max_files=1)
    ids = attachments.ConversationFileIds([1, '1'])
    ids.omitted_history_ids = [2, 3]
    text, snapshots = await attachments.load_context(ids, None, None, config)
    assert reads == [1]
    assert [row['id'] for row in snapshots] == [1]
    assert '2, 3' in text and '不得将其视为本轮已读取材料' in text
    reads.clear()
    with pytest.raises(ChatConfigurationError) as error:
        await attachments.load_context([1, 2], None, None, config)
    assert error.value.reason == 'attachment_budget_exceeded'
    assert reads == []


@pytest.mark.asyncio
@pytest.mark.parametrize('protection', ['other_session', 'message', 'skill', 'knowledge',
                                        'tenant', 'legacy', 'other_owner', 'same_path', 'exclusive'])
async def test_conversation_cleanup_preserves_shared_files(attachment_session, monkeypatch, tmp_path, protection):
    session = attachment_session
    principal = SimpleNamespace(tenant_id=1, user_id=2)
    source = tmp_path / 'retained.txt'
    source.write_bytes(b'retained fixture')
    data = {'accessMode': 'private', 'uploadToken': 'fixture-token',
            'chatSessionIds': ['10'], 'path': str(source)}
    if protection == 'other_session':
        data['chatSessionIds'].append('20')
    if protection in {'tenant', 'legacy'}:
        data.pop('accessMode')
        if protection == 'tenant':
            data['accessMode'] = 'tenant'
    file = Resource(id=1, tenant_id=1, owner_id=3 if protection == 'other_owner' else 2,
                    resource_type='files', data=data)
    session.add(file)
    add_attachment_message(session, 100, ['1'])
    if protection == 'message':
        add_attachment_message(session, 101, [1], conversation=20)
    elif protection == 'skill':
        session.add(Resource(id=102, tenant_id=1, owner_id=2, resource_type='skills',
                             data={'steps': [{'skillFile': {'id': '1'}}]}))
    elif protection == 'knowledge':
        session.add(KnowledgeDocument(id=103, tenant_id=1, knowledge_base_id=50,
                                      file_id=1, source_name='fixture'))
    elif protection == 'same_path':
        session.add(Resource(id=104, tenant_id=1, owner_id=2, resource_type='files',
                             data={'path': str(source), 'accessMode': 'tenant'}))
    await session.flush()
    monkeypatch.setattr(attachments, 'file_source_path', lambda file: source)
    paths = await attachments.cleanup_conversation(10, principal, session)
    assert (await session.get(Resource, 100)).is_deleted
    assert file.is_deleted is (protection in {'exclusive', 'same_path'})
    assert paths == ([source] if protection == 'exclusive' else [])
    assert source.read_bytes() == b'retained fixture'
    if protection == 'other_session':
        assert file.data['chatSessionIds'] == ['20']
    if protection == 'message':
        assert not (await session.get(Resource, 101)).is_deleted


def limits(**overrides):
    return SimpleNamespace(chat_attachment_max_bytes=1024 * 1024,
        chat_attachment_max_expanded_bytes=16 * 1024 * 1024,
        chat_attachment_max_pages=20, chat_attachment_max_chars=60000, **overrides)


@pytest.mark.parametrize('extension', ['txt', 'md', 'csv', 'json'])
def test_text_attachment_roundtrip(extension):
    assert '合同编号 123' in attachments.extract_payload('合同编号 123'.encode(), 'fixture.' + extension, limits())


def test_pdf_attachment_roundtrip():
    import pymupdf
    document = pymupdf.open()
    document.new_page().insert_text((40, 40), 'EIC Attachment Fixture 735')
    payload = document.tobytes()
    document.close()
    assert '735' in attachments.extract_payload(payload, 'fixture.pdf', limits())


def test_docx_attachment_roundtrip():
    from docx import Document
    document = Document()
    document.add_paragraph('文档编号 736')
    output = BytesIO()
    document.save(output)
    assert '736' in attachments.extract_payload(output.getvalue(), 'fixture.docx', limits())


def test_excel_attachment_roundtrip():
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.active.append(['编号', '金额'])
    workbook.active.append(['737', 18])
    output = BytesIO()
    workbook.save(output)
    assert '737' in attachments.extract_payload(output.getvalue(), 'fixture.xlsx', limits())


def test_presentation_attachment_roundtrip_and_entity_rejection():
    output = BytesIO()
    with zipfile.ZipFile(output, 'w') as archive:
        archive.writestr('ppt/slides/slide1.xml',
            '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            '<a:t>幻灯片 738</a:t></p:sld>')
    assert '738' in attachments.extract_payload(output.getvalue(), 'fixture.pptx', limits())
    malicious = BytesIO()
    with zipfile.ZipFile(malicious, 'w') as archive:
        archive.writestr('ppt/slides/slide1.xml', '<!DOCTYPE x [<!ENTITY x "secret">]><slide>&x;</slide>')
    with pytest.raises(ChatConfigurationError):
        attachments.extract_payload(malicious.getvalue(), 'fixture.pptx', limits())


@pytest.mark.parametrize('payload,name', [(b'', 'empty.txt'), (b'\x00', 'nul.txt'),
                                        (b'not-a-pdf', 'broken.pdf'), (b'code', 'run.exe')])
def test_invalid_attachment_rejected(payload, name):
    with pytest.raises(ChatConfigurationError):
        attachments.extract_payload(payload, name, limits())


def test_image_attachment_ocr_roundtrip():
    from PIL import Image, ImageDraw, ImageFont
    image = Image.new('RGB', (650, 120), 'white')
    ImageDraw.Draw(image).text((20, 25), 'Invoice 739', fill='black', font=ImageFont.load_default(size=42))
    output = BytesIO()
    image.save(output, format='PNG')
    assert '739' in attachments.extract_payload(output.getvalue(), 'fixture.png', limits())


@pytest.mark.asyncio
async def test_saved_batch_32_is_split_before_send(monkeypatch):
    calls = []
    def reply(request):
        rows = json.loads(request.content)['input']
        calls.append(len(rows))
        assert len(rows) <= 8
        return httpx.Response(200, json={'data':[{'index':i,'embedding':[float(text),1.0]}
            for i,text in reversed(list(enumerate(rows)))]})
    model = AIModel(id=1, tenant_id=1, model_type='embedding', provider='openai',
        model_name='fixture', endpoint='http://fixture.invalid/v1', enabled=True, embedding_dimensions=2)
    async with httpx.AsyncClient(transport=httpx.MockTransport(reply)) as client:
        vectors = await embeddings.embed_texts(model, [str(i) for i in range(65)],
                                              client=client, tenant_id=1, batch_size=32)
    assert calls == [8] * 8 + [1]
    assert [v[0] for v in vectors] == list(range(65))


@pytest.mark.asyncio
async def test_web_search_sources_filtered_and_bounded(monkeypatch):
    from eic.core import dependency_controls
    @asynccontextmanager
    async def guard(*args, **kwargs):
        yield
    monkeypatch.setattr(dependency_controls, 'guard', guard)
    original = httpx.AsyncClient
    def handler(request):
        assert request.url.params['format'] == 'json'
        return httpx.Response(200, json={'results':[
            {'url':'javascript:alert(1)','title':'invalid'},
            {'url':'https://example.org/reference','title':'Reference','content':'Verified fixture'},
            {'url':'https://example.org/reference','title':'duplicate'}]})
    monkeypatch.setattr(search.httpx, 'AsyncClient', lambda **kwargs:
        original(transport=httpx.MockTransport(handler), **kwargs))
    config = SimpleNamespace(chat_search_backend='searxng', chat_search_endpoint='http://fixture.invalid/search',
        chat_search_api_key=SecretStr(''), chat_search_timeout_seconds=1,
        chat_search_max_chars=1000, chat_search_max_results=5)
    context, citations = await search.independent_search('fixture', config, tenant_id=1)
    assert len(citations) == 1 and citations[0]['type'] == 'web'
    assert 'https://example.org/reference' in context


def test_search_mode_requires_explicit_configuration():
    config = SimpleNamespace(chat_search_backend='none',chat_search_endpoint='',chat_search_api_key=SecretStr(''))
    with pytest.raises(ChatConfigurationError):
        search.validate_mode(SimpleNamespace(search_mode='independent'), config)
    with pytest.raises(ChatConfigurationError):
        search.validate_mode(SimpleNamespace(search_mode='native', provider='anthropic'), config)
