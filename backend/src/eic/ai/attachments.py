"""会话附件：本地抽取、重新授权及累计预算；图片仅 OCR。"""
import asyncio
import hashlib
from io import BytesIO
from pathlib import Path
import re
import zipfile
from xml.etree import ElementTree

from sqlalchemy import select, or_, and_

from eic.ai.budget import ChatConfigurationError
from eic.core.file_access import file_source_path, require_file_access, valid_file_id
from eic.models import ChatMessage, Resource, KnowledgeDocument

FORMATS = frozenset({'pdf', 'docx', 'txt', 'md', 'json', 'png', 'jpg', 'jpeg', 'webp',
                     'bmp', 'tif', 'tiff', 'gif', 'xlsx', 'xlsm', 'xls', 'csv', 'pptx'})


def extract_payload(payload: bytes, filename: str, settings) -> str:
    """体验模块可直接复用；调用者负责文件授权和原始字节总预算。"""
    suffix = Path(filename).suffix.lower()
    if suffix.lstrip('.') not in FORMATS:
        raise ChatConfigurationError('attachment_format_unsupported')
    if len(payload) > settings.chat_attachment_max_bytes:
        raise ChatConfigurationError('attachment_budget_exceeded')
    try:
        if suffix in {'.txt', '.md', '.csv', '.json'}:
            text = payload.decode('utf-8-sig', errors='strict')
        else:
            if suffix in {'.docx', '.xlsx', '.xlsm', '.pptx'}:
                with zipfile.ZipFile(BytesIO(payload)) as archive:
                    entries = archive.infolist()
                    if (len(entries) > 5000 or sum(i.file_size for i in entries) > settings.chat_attachment_max_expanded_bytes
                            or any(i.flag_bits & 1 for i in entries)):
                        raise ChatConfigurationError('attachment_budget_exceeded')
                    if suffix == '.pptx':
                        slides = sorted((i for i in entries if re.fullmatch(r'ppt/slides/slide\d+\.xml', i.filename)),
                                        key=lambda i: int(re.search(r'slide(\d+)', i.filename).group(1)))
                        if len(slides) > settings.chat_attachment_max_pages:
                            raise ChatConfigurationError('attachment_budget_exceeded')
                        parts = []
                        for index, info in enumerate(slides, 1):
                            xml = archive.read(info)
                            if b'<!DOCTYPE' in xml.upper() or b'<!ENTITY' in xml.upper():
                                raise ValueError('unsafe XML')
                            root = ElementTree.fromstring(xml)
                            slide_text = '\n'.join(node.text or '' for node in root.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}t'))
                            if slide_text.strip():
                                parts.append(f'[幻灯片 {index}]\n' + slide_text)
                        text = '\n\n'.join(parts)
                        if len(text) > settings.chat_attachment_max_chars:
                            raise ChatConfigurationError('attachment_budget_exceeded')
                        if not text.strip():
                            raise ValueError('empty presentation')
                        return text
            # 解析器以扩展名分派；不依赖上传存储文件名，也不调用远端 sidecar。
            import tempfile
            from eic.knowledge.documents import extract_document
            with tempfile.TemporaryDirectory(prefix='eic-chat-') as directory:
                local = Path(directory) / ('attachment' + suffix)
                local.write_bytes(payload)
                if suffix == '.pdf':
                    import pymupdf
                    with pymupdf.open(local) as document:
                        if document.page_count > settings.chat_attachment_max_pages:
                            raise ChatConfigurationError('attachment_budget_exceeded')
                if suffix.lstrip('.') in {'png', 'jpg', 'jpeg', 'webp', 'bmp', 'tif', 'tiff', 'gif'}:
                    from PIL import Image
                    with Image.open(local) as image:
                        if image.width * image.height > 20000000:
                            raise ChatConfigurationError('attachment_budget_exceeded')
                if suffix in {'.xlsx', '.xlsm', '.xls'}:
                    text = _extract_excel(local, suffix, settings.chat_attachment_max_chars)
                else:
                    text = extract_document(local).text
        if not text.strip() or '\x00' in text:
            raise ValueError('empty document')
        if len(text) > settings.chat_attachment_max_chars:
            raise ChatConfigurationError('attachment_budget_exceeded')
        return text
    except ChatConfigurationError:
        raise
    except Exception as exc:
        raise ChatConfigurationError('attachment_parse_failed') from exc


def _extract_excel(path, suffix, max_chars):
    """逐行检查字符/单元格预算，避免稀疏工作表膨胀后再检查。"""
    parts, chars, cells = [], 0, 0
    if suffix == '.xls':
        import xlrd
        book = xlrd.open_workbook(path, on_demand=True)
        sheets = ((sheet.name, (tuple(cell.value for cell in row) for row in sheet.get_rows())) for sheet in book.sheets())
        close = book.release_resources
    else:
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        sheets = ((sheet.title, sheet.iter_rows(values_only=True)) for sheet in book.worksheets)
        close = book.close
    try:
        for name, rows in sheets:
            parts.append(f'[工作表 {name}]')
            chars += len(parts[-1]) + 1
            for row in rows:
                cells += len(row)
                if cells > 100000:
                    raise ChatConfigurationError('attachment_budget_exceeded')
                line = '\t'.join('' if value is None else str(value) for value in row)
                chars += len(line) + 1
                if chars > max_chars:
                    raise ChatConfigurationError('attachment_budget_exceeded')
                if line.strip():
                    parts.append(line)
        if chars > max_chars:
            raise ChatConfigurationError('attachment_budget_exceeded')
        return '\n'.join(parts)
    finally:
        close()


class ConversationFileIds(list):
    """保留列表接口，并将历史截断说明传给每次上下文重读。"""

    def __init__(self, values):
        super().__init__(values)
        self.omitted_history_ids = []
        self.history_truncated = False


async def conversation_file_ids(current_ids, principal, session, conversation_id=None, *, before_message_id=None,
                                max_files=None):
    """本轮附件全部保留，历史附件按最近消息优先补至数量预算。"""
    if max_files is None:
        from eic.core.config import get_settings
        max_files = get_settings().chat_attachment_max_files
    ids = ConversationFileIds(dict.fromkeys(valid_file_id(value) for value in (current_ids or [])))
    if len(ids) > max_files:
        raise ChatConfigurationError('attachment_budget_exceeded')
    if conversation_id is None:
        return ids
    query = select(Resource.data).join(ChatMessage, ChatMessage.id == Resource.id).where(
        Resource.tenant_id == principal.tenant_id, Resource.owner_id == principal.user_id,
        Resource.resource_type == 'chat_messages', Resource.is_deleted.is_(False),
        ChatMessage.tenant_id == principal.tenant_id, ChatMessage.session_id == conversation_id,
        ChatMessage.role == 'user')
    if before_message_id is not None:
        boundary = await session.get(ChatMessage, before_message_id)
        if boundary is not None:
            query = query.where(or_(ChatMessage.created_at < boundary.created_at,
                and_(ChatMessage.created_at == boundary.created_at, ChatMessage.id <= boundary.id)))
    rows = list(await session.scalars(query.order_by(ChatMessage.created_at.desc(), ChatMessage.id.desc()).limit(501)))
    ids.history_truncated = len(rows) > 500
    rows = rows[:500]
    seen = set(ids)
    for data in rows:
        for value in (data or {}).get('fileIds') or []:
            value = valid_file_id(value)
            if value in seen:
                continue
            seen.add(value)
            if len(ids) < max_files:
                ids.append(value)
            else:
                ids.omitted_history_ids.append(value)
    return ids


async def bind_to_conversation(ids, principal, session, conversation_id):
    for value in ids:
        file = await require_file_access(session, principal, value)
        data = file.data or {}
        # 仅认领可证明属于上传者的私有文件；历史和租户共享文件不转为聊天专属。
        if data.get('accessMode') == 'private' and data.get('uploadToken') and file.owner_id == principal.user_id:
            file = await require_file_access(session, principal, value, write=True)
            owners = list(data.get('chatSessionIds') or [])
            if str(conversation_id) not in owners:
                owners.append(str(conversation_id))
            file.data = {**data, 'chatSessionIds': owners}


def _references_file(data, file_id):
    if isinstance(data, dict):
        for key, value in data.items():
            if key in {'fileId', 'fileIds'}:
                values = value if isinstance(value, list) else [value]
                if str(file_id) in {str(v) for v in values}:
                    return True
            if key in {'file', 'skillFile'} and isinstance(value, dict) and str(value.get('id')) == str(file_id):
                return True
            if _references_file(value, file_id):
                return True
    elif isinstance(data, list):
        return any(_references_file(value, file_id) for value in data)
    return False


async def cleanup_conversation(conversation_id, principal, session):
    """事务内标记专属文件删除；返回提交后才可清理的物理路径。"""
    messages = list(await session.scalars(select(Resource).join(ChatMessage, ChatMessage.id == Resource.id).where(
        Resource.tenant_id == principal.tenant_id, Resource.owner_id == principal.user_id,
        Resource.resource_type == 'chat_messages', ChatMessage.session_id == conversation_id)))
    for row in messages:
        row.is_deleted = True
    await session.flush()
    resources = list(await session.scalars(select(Resource).where(Resource.tenant_id == principal.tenant_id,
        Resource.is_deleted.is_(False))))
    paths = []
    for file in resources:
        data = file.data or {}
        if (file.resource_type != 'files' or file.owner_id != principal.user_id
                or data.get('accessMode') != 'private' or not data.get('uploadToken')
                or str(conversation_id) not in data.get('chatSessionIds', [])):
            continue
        other_owners = [v for v in data['chatSessionIds'] if v != str(conversation_id)]
        file.data = {**data, 'chatSessionIds': other_owners}
        if other_owners or any(_references_file(row.data, file.id) for row in resources if row.id != file.id):
            continue
        shared = await session.scalar(select(KnowledgeDocument.id).where(
            KnowledgeDocument.tenant_id == principal.tenant_id, KnowledgeDocument.file_id == file.id).limit(1))
        if shared is not None:
            continue
        try:
            source = file_source_path(file)
        except ValueError:
            source = None
        file.is_deleted = True
        if source is not None and not any((row.data or {}).get('path') == data.get('path')
                for row in resources if row.id != file.id and not row.is_deleted):
            paths.append(source)
    return paths


async def purge_paths(paths):
    for path in paths:
        try:
            await asyncio.to_thread(path.unlink, missing_ok=True)
        except OSError:
            # 已撤销访问；保留清理失败的孤立字节供运维回收。
            import logging
            logging.getLogger(__name__).warning('聊天附件物理清理失败')


def enabled(settings) -> bool:
    return settings.chat_attachments_enabled and all(
        getattr(settings, name, None) is not None
        for name in ("chat_attachment_max_files", "chat_attachment_max_bytes", "chat_attachment_max_chars")
    )


def read_bounded(path, limit):
    with path.open("rb") as stream:
        content = stream.read(limit + 1)
    if len(content) > limit:
        raise ChatConfigurationError("attachment_budget_exceeded")
    return content


async def load_context(ids, principal, session, settings, *, parsed_cache=None):
    """缓存仅由单次请求持有；每次仍重新授权、读文件及校验原始内容。"""
    omitted_history_ids = getattr(ids, 'omitted_history_ids', [])
    history_truncated = getattr(ids, 'history_truncated', False)
    ids = list(dict.fromkeys(valid_file_id(value) for value in (ids or [])))
    scan_note = ('[附件回读范围说明]\n历史附件仅扫描最近 500 条用户消息，更早消息未扫描，'
                 '其中附件本轮未回读；不得视为已读取材料。') if history_truncated else ''
    if not ids:
        if scan_note:
            return '\n\n' + scan_note, []
        return "", []
    if not enabled(settings):
        raise ChatConfigurationError("attachment_context_not_supported")
    if len(ids) > settings.chat_attachment_max_files:
        raise ChatConfigurationError("attachment_budget_exceeded")
    byte_count = char_count = 0
    texts, snapshots = [], []
    for value in ids:
        file = await require_file_access(session, principal, value)
        if file.name.rsplit(".", 1)[-1].lower() not in FORMATS:
            raise ChatConfigurationError("attachment_format_unsupported")
        try:
            source = file_source_path(file)
            payload = await asyncio.to_thread(read_bounded, source, settings.chat_attachment_max_bytes - byte_count)
            digest = hashlib.sha256(payload).hexdigest()
            cached = parsed_cache.get(file.id) if parsed_cache is not None else None
            if cached is not None and cached[:2] == (digest, file.name):
                text = cached[2]
            else:
                text = await asyncio.to_thread(extract_payload, payload, file.name, settings)
        except ChatConfigurationError:
            raise
        except (OSError, ValueError, UnicodeError) as exc:
            raise ChatConfigurationError("attachment_parse_failed") from exc
        if not text.strip() or "\x00" in text:
            raise ChatConfigurationError("attachment_parse_failed")
        byte_count += len(payload)
        char_count += len(text)
        if char_count > settings.chat_attachment_max_chars:
            raise ChatConfigurationError("attachment_budget_exceeded")
        if parsed_cache is not None:
            # 每个文件只保留一个版本，且缓存总项数受附件数量限制。
            if file.id not in parsed_cache and len(parsed_cache) >= settings.chat_attachment_max_files:
                parsed_cache.clear()
            parsed_cache[file.id] = (digest, file.name, text)
        snapshots.append({"id": file.id, "name": file.name, "sha256": digest, "size": len(payload)})
        # 附件处于本轮 user 消息，不提权为 system 指令；分隔符不是安全隔离保证。
        texts.append(f"文件 ID {file.id}，SHA-256 {digest}，以下为用户提供的未信任材料：\n{text}")
    if omitted_history_ids:
        texts.append('[附件回读范围说明]\n本轮附件已全部纳入；历史附件按最近消息优先回读。'
                     f'受每轮最多 {settings.chat_attachment_max_files} 份附件限制，以下历史文件 ID 本轮未回读：'
                     + ', '.join(str(value) for value in omitted_history_ids)
                     + '。这些文件仍保留在会话中；不得将其视为本轮已读取材料。')
    if scan_note:
        texts.append(scan_note)
    return "\n\n[用户附件材料]\n" + "\n\n".join(texts), snapshots
