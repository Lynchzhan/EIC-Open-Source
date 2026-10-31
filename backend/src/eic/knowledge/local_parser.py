"""可终止的本地文档解析进程；只接收路径和资源限额，不接收凭据。"""
import asyncio
import importlib
import json
import math
import os
from pathlib import Path
import signal
import subprocess
import sys
import tempfile
import time
import warnings
import zipfile

from eic.knowledge.documents import ExtractedDocument

FIELDS = ("max_bytes", "max_expanded_bytes", "max_pages", "max_pixels", "max_cells",
          "max_chars", "parse_timeout_seconds", "parse_memory_mb")
ERROR_CODES = frozenset({
    "document_source_limit", "document_expansion_limit", "document_page_limit", "document_pixel_limit",
    "document_cell_limit", "document_output_limit", "document_parse_timeout", "document_memory_limit",
    "document_parse_failed_or_limit_exceeded",
})


class LocalParseError(ValueError):
    def __init__(self, code):
        self.code = code if code in ERROR_CODES else "document_parse_failed_or_limit_exceeded"
        super().__init__(self.code)


def preflight(path, limits):
    if path.stat().st_size > limits["max_bytes"]:
        raise ValueError("document_source_limit")
    if path.suffix.lower() in {".docx", ".xlsx", ".xlsm", ".pptx"}:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if (len(entries) > 5000 or sum(i.file_size for i in entries) > limits["max_expanded_bytes"]
                    or any(i.flag_bits & 1 for i in entries)):
                raise ValueError("document_expansion_limit")
    if path.suffix.lower() == ".pdf":
        import pymupdf
        with pymupdf.open(path) as document:
            if document.page_count > limits["max_pages"]:
                raise ValueError("document_page_limit")
            for page in document:
                pixels = page.rect.width * page.rect.height * 4
                if not math.isfinite(pixels) or pixels > limits["max_pixels"]:
                    raise ValueError("document_pixel_limit")
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            cells = 0
            for sheet in book.worksheets:
                cells += (sheet.max_row or 0) * (sheet.max_column or 0)
                if cells > limits["max_cells"]:
                    raise ValueError("document_cell_limit")
        finally:
            book.close()
    if path.suffix.lower() == ".xls":
        import xlrd
        book = xlrd.open_workbook(path, on_demand=True)
        try:
            if sum(sheet.nrows * sheet.ncols for sheet in book.sheets()) > limits["max_cells"]:
                raise ValueError("document_cell_limit")
        finally:
            book.release_resources()


def worker(request_path, output_path):
    request = json.loads(Path(request_path).read_text(encoding="utf-8"))
    limits = request["limits"]
    if os.name == "posix":
        import resource
        # 虚拟地址空间与父进程监控的 RSS 分开限制，给数值库保留映射空间。
        ceiling = max(2048, limits["parse_memory_mb"] * 3) * 1024 * 1024
        resource.setrlimit(resource.RLIMIT_AS, (ceiling, ceiling))
        cpu = limits["parse_timeout_seconds"]
        resource.setrlimit(resource.RLIMIT_CPU, (cpu, cpu + 1))
    from PIL import Image
    Image.MAX_IMAGE_PIXELS = limits["max_pixels"]
    warnings.simplefilter("error", Image.DecompressionBombWarning)
    path = Path(request["path"])
    preflight(path, limits)
    module, name = request["extractor"]
    target = importlib.import_module(module)
    for part in name.split("."):
        target = getattr(target, part)
    extracted = target(path)
    if len(extracted.text) > limits["max_chars"] or sum(len(text) for _, text in extracted.pages) > limits["max_chars"]:
        raise ValueError("document_output_limit")
    payload = json.dumps({"text": extracted.text, "report": extracted.report, "pages": extracted.pages}, ensure_ascii=False)
    if len(payload.encode("utf-8")) > limits["max_chars"] * 12 + 1024 * 1024:
        raise ValueError("document_output_limit")
    Path(output_path).write_text(payload, encoding="utf-8")


def resident_bytes(pid):
    if sys.platform.startswith("linux"):
        for line in Path(f"/proc/{pid}/status").read_text().splitlines():
            if line.startswith("VmRSS:"):
                return int(line.split()[1]) * 1024
        return 0
    if os.name == "nt":
        import ctypes
        from ctypes import wintypes
        class Counters(ctypes.Structure):
            _fields_ = [("cb", wintypes.DWORD), ("PageFaultCount", wintypes.DWORD)] + [
                (name, ctypes.c_size_t) for name in ("PeakWorkingSetSize", "WorkingSetSize", "QuotaPeakPagedPoolUsage",
                "QuotaPagedPoolUsage", "QuotaPeakNonPagedPoolUsage", "QuotaNonPagedPoolUsage", "PagefileUsage", "PeakPagefileUsage")]
        kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel.OpenProcess.argtypes = [wintypes.DWORD, wintypes.BOOL, wintypes.DWORD]
        kernel.OpenProcess.restype = wintypes.HANDLE
        kernel.CloseHandle.argtypes = [wintypes.HANDLE]
        kernel.CloseHandle.restype = wintypes.BOOL
        handle = kernel.OpenProcess(0x410, False, pid)
        if not handle:
            raise OSError("parser_process_unavailable")
        try:
            counters = Counters()
            counters.cb = ctypes.sizeof(counters)
            psapi = ctypes.WinDLL("psapi", use_last_error=True)
            psapi.GetProcessMemoryInfo.argtypes = [wintypes.HANDLE, ctypes.POINTER(Counters), wintypes.DWORD]
            psapi.GetProcessMemoryInfo.restype = wintypes.BOOL
            if not psapi.GetProcessMemoryInfo(handle, ctypes.byref(counters), counters.cb):
                raise OSError("parser_memory_unavailable")
            return counters.WorkingSetSize
        finally:
            kernel.CloseHandle(handle)
    raise OSError("parser_platform_unsupported")


async def extract_local(path, settings, extractor):
    if "<locals>" in extractor.__qualname__ or extractor.__name__ == "<lambda>":
        raise ValueError("local_extractor_requires_importable_function")
    limits = {name: getattr(settings, "document_" + name) for name in FIELDS}
    with tempfile.TemporaryDirectory(prefix="eic-parser-") as directory:
        request, output = Path(directory) / "request.json", Path(directory) / "result.json"
        request.write_text(json.dumps({"path": str(Path(path).resolve()), "limits": limits,
            "extractor": [extractor.__module__, extractor.__qualname__]}), encoding="utf-8")
        environment = {key: os.environ[key] for key in ("PATH", "SYSTEMROOT", "WINDIR", "TEMP", "TMP") if key in os.environ}
        environment.update(PYTHONPATH=str(Path(__file__).resolve().parents[2]),
                           OPENBLAS_NUM_THREADS="1", OMP_NUM_THREADS="1", PYTHONIOENCODING="utf-8")
        options = {"start_new_session": True} if os.name == "posix" else {"creationflags": subprocess.CREATE_NO_WINDOW}
        process = await asyncio.create_subprocess_exec(sys.executable, "-m", "eic.knowledge.local_parser", str(request), str(output),
            env=environment, stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, **options)
        waiter = asyncio.create_task(process.wait())
        deadline = time.monotonic() + limits["parse_timeout_seconds"]
        try:
            while not waiter.done():
                await asyncio.wait({waiter}, timeout=0.2)
                if waiter.done():
                    break
                if time.monotonic() >= deadline:
                    raise LocalParseError("document_parse_timeout")
                if resident_bytes(process.pid) > limits["parse_memory_mb"] * 1024 * 1024:
                    raise LocalParseError("document_memory_limit")
            if await waiter != 0 or not output.is_file():
                raise LocalParseError("document_parse_failed_or_limit_exceeded")
            if output.stat().st_size > limits["max_chars"] * 12 + 1024 * 1024:
                raise LocalParseError("document_output_limit")
            result = json.loads(await asyncio.to_thread(output.read_text, encoding="utf-8"))
            if "errorCode" in result:
                raise LocalParseError(result["errorCode"])
            return ExtractedDocument(result["text"], result["report"], tuple(tuple(page) for page in result["pages"]))
        finally:
            if process.returncode is None:
                if os.name == "posix":
                    try:
                        os.killpg(process.pid, signal.SIGKILL)
                    except ProcessLookupError:
                        pass
                else:
                    process.kill()
            await asyncio.shield(waiter)


if __name__ == "__main__":
    try:
        worker(*sys.argv[1:])
    except Exception as exc:
        code = str(exc) if isinstance(exc, ValueError) and str(exc) in ERROR_CODES else "document_parse_failed_or_limit_exceeded"
        Path(sys.argv[2]).write_text(json.dumps({"errorCode": code}), encoding="utf-8")
